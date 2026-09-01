# main.py
from fastapi import FastAPI, Depends, UploadFile, File, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm

from typing import List, Optional
from sqlalchemy.orm import Session
from datetime import date, datetime, timedelta
from io import BytesIO

from jose import JWTError, jwt
import hashlib
import hmac
import os

from openpyxl import load_workbook

from database import init_db, get_db
from models import Asset, AssetEvent, User
from schemas import (
    AssetCreate,
    AssetRead,
    AssetUpdate,
    AssetEventCreate,
    AssetEventRead,
    UserCreate,
    UserRead,
    Token,
    TokenData,
    PasswordChange,
    PasswordReset,
)

# ==========================
# CONFIG
# ==========================
LIFESPAN_YEARS_DEFAULT = 5  # RSR policy: 5-year asset lifespan

# ⚠️ For production: move to environment variables
SECRET_KEY = os.getenv("ASSETDESK_SECRET_KEY", "rsr-assetdesk-change-me-to-a-strong-random-secret")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60  # 1 hour

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="token")


# ==========================
# PASSWORD HASHING (PBKDF2)
# ==========================
def hash_password(password: str) -> str:
    """
    Hash password using PBKDF2-HMAC (sha256) + random salt.
    Stored format: "<salt_hex>$<hash_hex>"
    """
    if not isinstance(password, str):
        raise ValueError("Password must be a string")

    salt = os.urandom(16)
    salt_hex = salt.hex()

    dk = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 100_000)
    hash_hex = dk.hex()

    return f"{salt_hex}${hash_hex}"


def verify_password(plain_password: str, stored_value: str) -> bool:
    """Verify a plain password against stored "salt$hash"."""
    try:
        salt_hex, hash_hex = stored_value.split("$", 1)
    except ValueError:
        return False

    salt = bytes.fromhex(salt_hex)
    expected_hash = bytes.fromhex(hash_hex)

    dk = hashlib.pbkdf2_hmac("sha256", plain_password.encode("utf-8"), salt, 100_000)
    return hmac.compare_digest(dk, expected_hash)


# ==========================
# JWT HELPERS
# ==========================
def create_access_token(data: dict, expires_delta: Optional[timedelta] = None) -> str:
    to_encode = data.copy()
    expire = datetime.utcnow() + (expires_delta or timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES))
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)


def get_user_by_username(db: Session, username: str) -> Optional[User]:
    return db.query(User).filter(User.username == username).first()


def authenticate_user(db: Session, username: str, password: str) -> Optional[User]:
    user = get_user_by_username(db, username)
    if not user:
        return None
    if not verify_password(password, user.hashed_password):
        return None
    return user


async def get_current_user(
    token: str = Depends(oauth2_scheme),
    db: Session = Depends(get_db),
) -> User:
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )

    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception

    user = get_user_by_username(db, username=username)
    if user is None:
        raise credentials_exception

    return user


async def get_current_active_user(current_user: User = Depends(get_current_user)) -> User:
    if not current_user.is_active:
        raise HTTPException(status_code=400, detail="Inactive user")
    return current_user


def require_admin(current_user: User = Depends(get_current_active_user)) -> User:
    if (current_user.role or "").strip() != "AssetDeskAdmin":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Not enough permissions. Admin role required.",
        )
    return current_user


# ==========================
# LIFESPAN CALCULATION
# ==========================
def calculate_life_remaining(
    purchase_date: Optional[date],
    lifespan_years: Optional[int] = None,
) -> Optional[int]:
    """Returns an integer 0–100, or None if purchase_date is missing."""
    if purchase_date is None:
        return None

    lifespan_years = lifespan_years or LIFESPAN_YEARS_DEFAULT

    today = date.today()
    total_days = lifespan_years * 365  # approximation
    age_days = (today - purchase_date).days

    if age_days <= 0:
        return 100
    if age_days >= total_days:
        return 0

    remaining_ratio = 1 - (age_days / total_days)
    return int(round(remaining_ratio * 100))


# ==========================
# FASTAPI APP
# ==========================
app = FastAPI(title="RSR Asset Desk Backend", version="1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:5174",
        "http://127.0.0.1:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def on_startup():
    init_db()


@app.get("/")
def root():
    return {"message": "RSR Asset Desk backend ready. Try /health or /docs."}


@app.get("/health")
def health_check():
    return {"status": "ok", "message": "RSR Asset Desk backend is running"}


# ======================
# AUTH / LOGIN (TOKEN)
# ======================
@app.post("/token", response_model=Token)
async def login_for_access_token(
    form_data: OAuth2PasswordRequestForm = Depends(),
    db: Session = Depends(get_db),
):
    user = authenticate_user(db, form_data.username, form_data.password)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect username or password",
            headers={"WWW-Authenticate": "Bearer"},
        )

    expires = timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    access_token = create_access_token(data={"sub": user.username}, expires_delta=expires)

    return {"access_token": access_token, "token_type": "bearer"}


# ======================
# USERS
# ======================

@app.post("/users", response_model=UserRead, status_code=status.HTTP_201_CREATED)
def bootstrap_create_first_user(user_in: UserCreate, db: Session = Depends(get_db)):
    """
    Bootstrap-only:
    - If there are ZERO users, allow creating the very first user.
    - First user becomes AssetDeskAdmin by default (if role not provided).
    """
    user_count = db.query(User).count()
    if user_count > 0:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User creation is restricted. Use /users/admin-create with an admin token.",
        )

    existing = db.query(User).filter(User.username == user_in.username).first()
    if existing:
        raise HTTPException(status_code=400, detail="Username already registered")

    hashed_pw = hash_password(user_in.password)
    role = user_in.role or "AssetDeskAdmin"

    user = User(
        username=user_in.username,
        full_name=user_in.full_name,
        email=user_in.email,
        role=role,
        hashed_password=hashed_pw,
        is_active=user_in.is_active if user_in.is_active is not None else True,
    )

    db.add(user)
    db.commit()
    db.refresh(user)
    return user


@app.post("/users/admin-create", response_model=UserRead, status_code=status.HTTP_201_CREATED)
def admin_create_user(
    user_in: UserCreate,
    db: Session = Depends(get_db),
    current_admin: User = Depends(require_admin),
):
    """
    Admin-only user creation (this is what your frontend should use).
    """
    existing = db.query(User).filter(User.username == user_in.username).first()
    if existing:
        raise HTTPException(status_code=400, detail="Username already registered")

    hashed_pw = hash_password(user_in.password)

    user = User(
        username=user_in.username,
        full_name=user_in.full_name,
        email=user_in.email,
        role=user_in.role or "AssetDeskUser",
        hashed_password=hashed_pw,
        is_active=user_in.is_active if user_in.is_active is not None else True,
    )

    db.add(user)
    db.commit()
    db.refresh(user)
    return user


@app.get("/users/me", response_model=UserRead)
async def read_users_me(current_user: User = Depends(get_current_active_user)):
    return current_user


@app.get("/users", response_model=List[UserRead])
def list_users(db: Session = Depends(get_db), current_admin: User = Depends(require_admin)):
    return db.query(User).order_by(User.username.asc()).all()


@app.delete("/users/{user_id}", status_code=status.HTTP_200_OK)
def disable_user(user_id: int, db: Session = Depends(get_db), current_admin: User = Depends(require_admin)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if current_admin.id == user.id:
        raise HTTPException(status_code=400, detail="You cannot disable your own account.")

    if not user.is_active:
        return {"message": "User already disabled", "user_id": user.id}

    user.is_active = False
    db.commit()
    db.refresh(user)

    return {
        "message": f"User {user.username} disabled.",
        "user_id": user.id,
        "username": user.username,
        "is_active": user.is_active,
    }

@app.delete("/users/{user_id}/hard", status_code=status.HTTP_200_OK)
def hard_delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    current_admin: User = Depends(require_admin),
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if current_admin.id == user.id:
        raise HTTPException(status_code=400, detail="You cannot delete your own account.")

    db.delete(user)
    db.commit()
    return {"message": f"User {user.username} permanently deleted.", "user_id": user_id}


@app.patch("/users/{user_id}/enable", status_code=status.HTTP_200_OK)
def enable_user(user_id: int, db: Session = Depends(get_db), current_admin: User = Depends(require_admin)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if current_admin.id == user.id:
        raise HTTPException(status_code=400, detail="You cannot modify your own account status.")

    if user.is_active:
        return {"message": "User already active", "user_id": user.id}

    user.is_active = True
    db.commit()
    db.refresh(user)

    return {
        "message": f"User {user.username} enabled.",
        "user_id": user.id,
        "username": user.username,
        "is_active": user.is_active,
    }


@app.post("/users/me/change-password", status_code=status.HTTP_200_OK)
def change_my_password(
    payload: PasswordChange,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    if not verify_password(payload.current_password, current_user.hashed_password):
        raise HTTPException(status_code=400, detail="Current password is incorrect")

    if len(payload.new_password) < 8:
        raise HTTPException(status_code=400, detail="New password must be at least 8 characters")

    current_user.hashed_password = hash_password(payload.new_password)
    db.commit()

    return {"message": "Password updated successfully"}


@app.post("/users/{user_id}/reset-password", status_code=status.HTTP_200_OK)
def admin_reset_password(
    user_id: int,
    payload: PasswordReset,
    db: Session = Depends(get_db),
    current_admin: User = Depends(require_admin),
):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if len(payload.new_password) < 8:
        raise HTTPException(status_code=400, detail="New password must be at least 8 characters")

    user.hashed_password = hash_password(payload.new_password)
    db.commit()

    return {"message": f"Password reset for {user.username}"}


# ======================
# ASSETS
# ======================
@app.get("/assets", response_model=List[AssetRead])
def list_assets(db: Session = Depends(get_db)):
    assets = db.query(Asset).all()
    for asset in assets:
        asset.life_remaining_percent = calculate_life_remaining(
            asset.purchase_date,
            asset.useful_life_years or LIFESPAN_YEARS_DEFAULT,
        )
    return assets


@app.post("/assets", response_model=AssetRead, dependencies=[Depends(require_admin)])
def create_asset(asset_in: AssetCreate, db: Session = Depends(get_db)):
    if asset_in.asset_tag:
        existing_tag = db.query(Asset).filter(Asset.asset_tag == asset_in.asset_tag).first()
        if existing_tag:
            raise HTTPException(status_code=400, detail="An asset with this Asset ID already exists.")

    if asset_in.serial_number:
        existing_serial = db.query(Asset).filter(Asset.serial_number == asset_in.serial_number).first()
        if existing_serial:
            raise HTTPException(status_code=400, detail="An asset with this Serial Number already exists.")

    lifespan = asset_in.useful_life_years or LIFESPAN_YEARS_DEFAULT

    asset = Asset(
        asset_tag=asset_in.asset_tag,
        category=asset_in.category,
        brand=asset_in.brand,
        model=asset_in.model,
        serial_number=asset_in.serial_number,
        hostname=asset_in.hostname,
        status=asset_in.status or "In Use",
        assigned_to=asset_in.assigned_to,
        department=asset_in.department,
        location=asset_in.location,
        purchase_date=asset_in.purchase_date,
        warranty_expiry=asset_in.warranty_expiry,
        cost=asset_in.cost,
        useful_life_years=lifespan,
        notes=asset_in.notes,
    )

    asset.life_remaining_percent = calculate_life_remaining(asset.purchase_date, lifespan)

    db.add(asset)
    db.commit()
    db.refresh(asset)
    return asset


@app.put("/assets/{asset_id}", response_model=AssetRead, dependencies=[Depends(require_admin)])
def update_asset(asset_id: int, asset_in: AssetUpdate, db: Session = Depends(get_db)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    if asset_in.asset_tag and asset_in.asset_tag != asset.asset_tag:
        existing_tag = db.query(Asset).filter(Asset.asset_tag == asset_in.asset_tag).first()
        if existing_tag:
            raise HTTPException(status_code=400, detail="An asset with this Asset ID already exists.")

    if asset_in.serial_number and asset_in.serial_number != asset.serial_number:
        existing_serial = db.query(Asset).filter(Asset.serial_number == asset_in.serial_number).first()
        if existing_serial:
            raise HTTPException(status_code=400, detail="An asset with this Serial Number already exists.")

    update_data = asset_in.dict(exclude_unset=True)
    for field, value in update_data.items():
        setattr(asset, field, value)

    lifespan = asset.useful_life_years or LIFESPAN_YEARS_DEFAULT
    asset.life_remaining_percent = calculate_life_remaining(asset.purchase_date, lifespan)

    db.commit()
    db.refresh(asset)
    return asset


@app.delete("/assets/{asset_id}", dependencies=[Depends(require_admin)])
def delete_asset(asset_id: int, db: Session = Depends(get_db)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    db.delete(asset)
    db.commit()
    return {"message": "Asset deleted successfully", "id": asset_id}


# ======================
# ASSET HISTORY (EVENTS)
# ======================
@app.post("/assets/{asset_id}/events", response_model=AssetEventRead, status_code=status.HTTP_201_CREATED)
def create_asset_event(asset_id: int, ev: AssetEventCreate, db: Session = Depends(get_db)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    event = AssetEvent(
        asset_id=asset_id,
        event_type=ev.event_type,
        actor=ev.actor,
        timestamp=ev.timestamp or datetime.utcnow(),
        from_status=ev.from_status,
        to_status=ev.to_status,
        details=ev.details,
    )

    db.add(event)
    db.commit()
    db.refresh(event)
    return event


@app.get("/assets/{asset_id}/events", response_model=List[AssetEventRead])
def list_asset_events(asset_id: int, db: Session = Depends(get_db)):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    return (
        db.query(AssetEvent)
        .filter(AssetEvent.asset_id == asset_id)
        .order_by(AssetEvent.timestamp.desc(), AssetEvent.id.desc())
        .all()
    )


# ======================
# EXCEL IMPORT
# ======================
@app.post("/import-excel", dependencies=[Depends(require_admin)])
async def import_assets_from_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Please upload an Excel file (.xlsx or .xls).")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not read Excel file: {e}")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Excel file is empty.")

    header_row = rows[0]
    data_rows = rows[1:]

    def norm_header(h):
        if h is None:
            return ""
        return (
            str(h).strip().lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
            .replace("(", "")
            .replace(")", "")
        )

    header_map = {norm_header(h): idx for idx, h in enumerate(header_row)}

    def col_idx(*keys):
        for k in keys:
            if k in header_map:
                return header_map[k]
        return None

    idx_asset_id = col_idx("assetid", "id", "assetcode")
    idx_desc = col_idx("description", "desc", "assetdescription")
    idx_user = col_idx("user", "owner", "allocatedto")
    idx_dept = col_idx("department", "dept")
    idx_loc = col_idx("location", "loc", "site")
    idx_type = col_idx("type", "assettype", "category")
    idx_serial = col_idx("serial", "serialnumber", "sn")
    idx_value = col_idx("valuezar", "value", "cost", "amount")
    idx_purchased = col_idx("purchasedate", "purchased", "purchase", "date")

    if idx_asset_id is None:
        raise HTTPException(status_code=400, detail="Excel must have an 'Asset ID' column (or similar).")

    existing_assets = db.query(Asset).all()
    existing_ids = {(a.asset_tag or "").strip().lower() for a in existing_assets if a.asset_tag}
    existing_serials = {(a.serial_number or "").strip().lower() for a in existing_assets if a.serial_number}

    created = 0
    skipped_duplicates = 0
    skipped_missing_id = 0
    total_rows = len(data_rows)

    for row in data_rows:
        if row is None:
            continue

        def get_val(idx):
            if idx is None or idx >= len(row):
                return None
            v = row[idx]
            return v.strip() if isinstance(v, str) else v

        asset_id_raw = get_val(idx_asset_id)
        if not asset_id_raw:
            skipped_missing_id += 1
            continue

        asset_id = str(asset_id_raw).strip()
        asset_id_key = asset_id.lower()

        if asset_id_key in existing_ids:
            skipped_duplicates += 1
            continue

        desc = get_val(idx_desc)
        user = get_val(idx_user)
        dept = get_val(idx_dept)
        loc = get_val(idx_loc)
        type_ = get_val(idx_type)
        serial = get_val(idx_serial)
        value_raw = get_val(idx_value)
        purchase_raw = get_val(idx_purchased)

        serial_key = (str(serial).strip().lower() if serial else None)
        if serial_key and serial_key in existing_serials:
            skipped_duplicates += 1
            continue

        cost = None
        if isinstance(value_raw, (int, float)):
            cost = float(value_raw)
        elif isinstance(value_raw, str) and value_raw.strip():
            try:
                cost = float(value_raw.replace(",", ""))
            except ValueError:
                cost = None

        purchase_date = None
        if isinstance(purchase_raw, datetime):
            purchase_date = purchase_raw.date()
        elif isinstance(purchase_raw, date):
            purchase_date = purchase_raw
        elif isinstance(purchase_raw, str) and purchase_raw.strip():
            txt = purchase_raw.strip()
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
                try:
                    purchase_date = datetime.strptime(txt, fmt).date()
                    break
                except ValueError:
                    continue

        brand = None
        model = None
        if desc:
            parts = str(desc).split(" ", 1)
            if len(parts) == 1:
                brand = parts[0]
                model = parts[0]
            else:
                brand, model = parts[0], parts[1]

        lifespan = LIFESPAN_YEARS_DEFAULT

        asset = Asset(
            asset_tag=asset_id,
            category=str(type_).strip() if type_ else "Other",
            brand=brand,
            model=model or desc,
            serial_number=str(serial).strip() if serial else None,
            hostname=None,
            status="In Use",
            assigned_to=str(user).strip() if user else None,
            department=str(dept).strip() if dept else None,
            location=str(loc).strip() if loc else None,
            purchase_date=purchase_date,
            warranty_expiry=None,
            cost=cost,
            useful_life_years=lifespan,
            notes=None,
        )

        asset.life_remaining_percent = calculate_life_remaining(asset.purchase_date, lifespan)

        db.add(asset)
        created += 1

        existing_ids.add(asset_id_key)
        if serial_key:
            existing_serials.add(serial_key)

    db.commit()

    return {
        "file": file.filename,
        "total_rows_in_sheet": total_rows,
        "imported": created,
        "skipped_duplicates": skipped_duplicates,
        "skipped_missing_asset_id": skipped_missing_id,
    }
